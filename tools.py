import os
import subprocess
import shutil
import sys
from pathlib import Path
from duckduckgo_search import DDGS
import re
import boto3

BASE_PROJECT_PATH = Path("./projects")


def create_directory(rel_path, base_path=BASE_PROJECT_PATH):
    full_path = Path(base_path) / rel_path
    full_path.mkdir(parents=True, exist_ok=True)
    return f"Directory created: {rel_path}"

def write_file(rel_path, content, base_path=BASE_PROJECT_PATH):
    full_path = Path(base_path) / rel_path
    # Ensure parent directories exist
    full_path.parent.mkdir(parents=True, exist_ok=True)
    with open(full_path, "w", encoding="utf-8") as f:
        f.write(content)
    return f"File written: {rel_path}"

def check_dangerous_command(command):
    """Detects destructive commands to prevent system damage."""
    dangerous_patterns = [
        r"rm\s+-rf\s+/", r"format\s+", r"mkfs", r"dd\s+if=", 
        r"del\s+/s\s+/q\s+C:", r"rd\s+/s\s+/q\s+C:"
    ]
    for pattern in dangerous_patterns:
        if re.search(pattern, command, re.IGNORECASE):
            return True
    return False

def run_command(command, timeout=30, cwd=None):
    """Executes shell commands with a security gate."""
    if check_dangerous_command(command):
        return "SECURITY ALERT: Destructive command detected. Operation blocked for safety."
    try:
        execution_cwd = cwd or BASE_PROJECT_PATH
        result = subprocess.run(
            command,
            shell=True,
            cwd=execution_cwd,
            capture_output=True,
            text=True,
            timeout=timeout
        )
        if result.returncode == 0:
            return f"Success:\n{result.stdout}"
        else:
            return f"Error (Exit Code {result.returncode}):\n{result.stderr}"
    except Exception as e:
        return f"Exception: {str(e)}"

def setup_environment(path):
    full_path = BASE_PROJECT_PATH / path
    
    # Python detection
    if (full_path / "requirements.txt").exists():
        print(f"[SETUP] Python project detected in {path}. Creating venv...")
        run_command(f"python -m venv {path}/venv")
        return "Python venv created."
    
    # Node detection
    if (full_path / "package.json").exists():
        print(f"[SETUP] Node.js project detected in {path}. Running npm install...")
        run_in_env(path, "npm install")
        return "Node.js dependencies installed."

    # Java/Spring Boot detection
    if (full_path / "pom.xml").exists():
        print(f"[SETUP] Java/Maven project detected in {path}. Ready for mvnw.")
        return "Java/Maven project detected."
    
    return "No specific environment needed or detected."

def run_in_env(project_rel_path, command):
    project_path = (BASE_PROJECT_PATH / project_rel_path).resolve()
    
    # Cross-platform venv detection
    venv_python_win = project_path / "venv" / "Scripts" / "python.exe"
    venv_python_lin = project_path / "venv" / "bin" / "python"
    
    venv_python = venv_python_win if venv_python_win.exists() else venv_python_lin
    
    if venv_python.exists():
        if command.startswith("pip "):
            full_cmd = f'"{venv_python}" -m ' + command
        elif command.startswith("python "):
            full_cmd = command.replace("python ", f'"{venv_python}" ')
        else:
            full_cmd = command
        return run_command(full_cmd, cwd=project_path)
    
    # Default to direct execution (Node.js, Go, etc.)
    return run_command(command, cwd=project_path)

def web_search(query):
    print(f"Searching web for: {query}")
    try:
        with DDGS() as ddgs:
            results = [r for r in ddgs.text(query, max_results=5)]
            return str(results)
    except Exception as e:
        return f"Search Error: {e}"

def delete_directory(rel_path):
    full_path = BASE_PROJECT_PATH / rel_path
    if full_path.exists() and full_path.is_dir():
        print(f"Cleaning up directory: {full_path}")
        shutil.rmtree(full_path, ignore_errors=True)
        return f"Directory {rel_path} deleted successfully."
    return f"Error: Directory {rel_path} not found."

def list_files(rel_path="."):
    full_path = BASE_PROJECT_PATH / rel_path
    files = []
    if not full_path.exists(): return []
    for f in full_path.rglob("*"):
        if f.is_file():
            files.append(str(f.relative_to(BASE_PROJECT_PATH)))
    return files

def kaggle_search(query):
    print(f"Searching Kaggle for: {query}")
    try:
        import os
        from kaggle.api.kaggle_api_extended import KaggleApi
        
        # Set environment variables for Kaggle API if present in our config
        from config import config
        if hasattr(config, 'KAGGLE_USERNAME') and config.KAGGLE_USERNAME:
            os.environ['KAGGLE_USERNAME'] = config.KAGGLE_USERNAME
            os.environ['KAGGLE_KEY'] = config.KAGGLE_KEY

        api = KaggleApi()
        api.authenticate()
        datasets = api.dataset_list(search=query)
        results = []
        for d in datasets[:3]:
            # Use getattr to be safe across versions
            ref = getattr(d, 'ref', 'unknown')
            title = getattr(d, 'title', 'unknown')
            results.append({"ref": ref, "title": title})
        return str(results)
    except Exception as e:
        return f"Kaggle Search Error: {e} (Make sure kaggle.json is in ~/.kaggle/)"

def kaggle_download(dataset_ref, path="."):
    full_path = BASE_PROJECT_PATH / path
    print(f"Downloading Kaggle dataset {dataset_ref} to {full_path}")
    try:
        import os
        from kaggle.api.kaggle_api_extended import KaggleApi
        
        # Set environment variables for Kaggle API if present in our config
        from config import config
        if hasattr(config, 'KAGGLE_USERNAME') and config.KAGGLE_USERNAME:
            os.environ['KAGGLE_USERNAME'] = config.KAGGLE_USERNAME
            os.environ['KAGGLE_KEY'] = config.KAGGLE_KEY

        api = KaggleApi()
        api.authenticate()
        api.dataset_download_files(dataset_ref, path=str(full_path), unzip=True)
        return f"Dataset {dataset_ref} downloaded and unzipped successfully."
    except Exception as e:
        return f"Kaggle Download Error: {e}"

def extract_text_from_pdf(file_path):
    import PyPDF2
    text = ""
    try:
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text += page.extract_text() + "\n"
        return text
    except Exception as e:
        return f"PDF Error: {e}"

def extract_text_from_pptx(file_path):
    from pptx import Presentation
    text = ""
    try:
        prs = Presentation(file_path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
        return text
    except Exception as e:
        return f"PPTX Error: {e}"

def zip_directory(rel_path, base_path=BASE_PROJECT_PATH):
    full_path = Path(base_path) / rel_path
    zip_path = Path(base_path) / f"{rel_path}.zip"
    if full_path.exists():
        shutil.make_archive(str(Path(base_path) / rel_path), 'zip', full_path)
        return f"Project zipped successfully: {rel_path}.zip"
    return "Error: Directory not found."

def upload_to_cloud(file_path, bucket_name="infinity-projects"):
    """Uploads a file to S3-compatible storage (SaaS ready)."""
    try:
        s3 = boto3.client('s3', 
                          endpoint_url=os.getenv("S3_ENDPOINT"),
                          aws_access_key_id=os.getenv("S3_KEY"),
                          aws_secret_access_key=os.getenv("S3_SECRET"))
        filename = os.path.basename(file_path)
        s3.upload_file(file_path, bucket_name, filename)
        return f"SUCCESS: File {filename} uploaded to cloud storage."
    except Exception as e:
        return f"CLOUD ERROR: {str(e)}"
